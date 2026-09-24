"""Build docs/judgement-calls.xlsx from wiki_xref's classification."""
import sys, collections
from pathlib import Path
import xml.etree.ElementTree as ET

REPO = Path(r"C:\src\WHFB6eBattlescribe")
sys.path.insert(0, str(REPO / ".claude/skills/bsdata-cat/scripts"))
from wiki_xref import (Wiki, analyse, cat_inventory, normalise, MECHANICAL_TAG,
                       STAT_ORDER, _best, ARMY_ASSOC)
from bsdata import source_files, CAT_NS, GST_NS

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

W = Wiki(Path("C:/src/6eWikiScraper/data"))
WIKI = "https://6th.whfb.app"

# Decisions already taken. A settled row is dropped from the sheet; one marked
# "Needs more thought" is carried forward with its note, so the work done so far
# is never asked for twice.
import json
DEC = {}
_rec = REPO / "docs/wiki-decisions.json"
if _rec.exists():
    for d in json.loads(_rec.read_text(encoding="utf-8"))["decisions"]:
        DEC[(d["file"], d["category"], d["name"], d["used_by"])] = d


def context_map(path):
    """id -> where it sits: the owning unit chain, or who links to a shared rule."""
    ns = "{%s}" % (GST_NS if path.suffix == ".gst" else CAT_NS)
    root = ET.fromstring(path.read_text(encoding="utf-8", newline=""))
    par = {c: p for p in root.iter() for c in p}

    def chain(el):
        out, cur = [], par.get(el)
        while cur is not None:
            t = cur.tag.split("}")[-1]
            if t in ("selectionEntry", "selectionEntryGroup") and cur.get("name"):
                out.append(cur.get("name"))
            cur = par.get(cur)
        return out          # innermost first

    # who infoLinks each shared rule
    linkers = collections.defaultdict(set)
    for il in root.iter(ns + "infoLink"):
        tid = il.get("targetId")
        if not tid:
            continue
        c = chain(il)
        linkers[tid].add(c[-1] if c else "(catalogue root)")

    out = {}
    for el in root.iter():
        eid = el.get("id")
        if not eid:
            continue
        c = chain(el)
        if c:
            # outermost first reads like a path: Plaguebearers > Plaguebearer
            out[eid] = " > ".join(reversed(c))
        elif eid in linkers:
            out[eid] = "shared; used by " + ", ".join(sorted(linkers[eid]))
        else:
            out[eid] = "shared; nothing links to it"
    return out


rows = []      # (category, name, where, file, issue, optA, optB, link)

for p in source_files(REPO):
    ctx = context_map(p)
    inv, rs = analyse(p, W)

    for r in rs:
        where = ctx.get(r.get("id"), "-")
        if r["bucket"] == "C":
            alts = [h["path"] for h in r.get("hits", [])]
            rows.append(("Ambiguous match", r["name"], where, p.name,
                         r.get("note") or "wiki has {} candidate pages".format(len(alts)),
                         alts[0] if alts else "", alts[1] if len(alts) > 1 else "",
                         WIKI + alts[0] if alts else ""))
        elif r["bucket"] == "E" and not MECHANICAL_TAG.match(r["name"].strip()):
            cands = r.get("candidates", [])
            best = cands[0] if cands else None
            rows.append(("No wiki coverage", r["name"], where, p.name,
                         "no confident match; nearest candidate is not applied",
                         "leave as is",
                         "{} ({:.2f})".format(best[0], best[2]) if best else "no candidate",
                         WIKI + best[0] if best else ""))

    for c in inv["costs"]:
        wp = _best(W, c["name"], "magic-item")
        if not wp:
            continue
        want = (W.pages[wp] or {}).get("points_cost")
        try:
            have = float(c["pts"])
        except (TypeError, ValueError):
            continue
        if want is None or abs(have - float(want)) < 0.001:
            continue
        rows.append(("Points cost", c["name"], ctx.get(c.get("id"), "-"), p.name,
                     "catalogue and wiki disagree on cost",
                     "catalogue: {:g}".format(have), "wiki: {:g}".format(float(want)),
                     WIKI + wp))

    blank = {"", "-", "--", "n/a"}
    for pr in inv["profiles"]:
        wp = _best(W, pr["name"], "unit")
        if not wp:
            continue
        stats = (W.pages[wp] or {}).get("statistics") or {}
        diff = []
        for k in STAT_ORDER:
            h = (pr["chars"].get(k) or "").strip()
            v = str(stats.get(k, "")).strip()
            if h.lower() in blank or v.lower() in blank or h == v:
                continue
            diff.append("{} {}->{}".format(k, h, v))
        if diff:
            rows.append(("Unit statline", pr["name"], ctx.get(pr.get("id"), "-"), p.name,
                         "{} characteristic(s) differ".format(len(diff)),
                         "catalogue as is", "wiki: " + ", ".join(diff), WIKI + wp))

carried = []
kept = []
for r in rows:
    d = DEC.get((r[3], r[0], r[1], r[2]))
    if d is None:
        kept.append(r + (None,))
    elif d["decision"] == "Needs more thought":
        carried.append(r + (d,))
    # anything else is settled and drops out
rows = kept + carried

order = {"Ambiguous match": 0, "No wiki coverage": 1, "Points cost": 2, "Unit statline": 3}
rows.sort(key=lambda r: (order[r[0]], r[3], r[1]))

wb = Workbook()
ARIAL = "Arial"

lg = wb.active
lg.title = "Legend"
lg["A1"] = "WHFB 6e - judgement calls needing a decision"
lg["A1"].font = Font(name=ARIAL, size=14, bold=True)
notes = [
    "",
    "Fill in the Decision column on the 'Decisions' sheet, or just write what should",
    "happen in Notes - a note on its own counts as an answer and will be actioned.",
    "Decision cells are shaded yellow and offer a dropdown; you may also type free text.",
    "Use the Notes column for anything an agent actioning this should know.",
    "",
    "'Used by' is what tells two same-named rules apart. For a rule that sits inside a",
    "unit it reads as a path, outermost first: 'Plaguebearers > Plaguebearer'. For one",
    "in <sharedRules> it lists every entry that links to it - so a rule showing",
    "'shared; used by Plaguebearers, Stream of Corruption' serves both the unit and the",
    "Daemonic Gift of the same name, which is the decision to make.",
    "",
    "Category meanings:",
    "  Ambiguous match   the wiki has more than one page for this name and the texts differ",
    "  No wiki coverage  no confident match. Often a catalogue spelling error - check the candidate",
    "  Points cost       catalogue and wiki disagree. NB Bretonnian Virtues are priced per",
    "                    character rank, so several rows per item is expected, not a bug",
    "  Unit statline     characteristics differ. Only characteristics both sides state are compared",
    "",
    "Source: wiki scrape at C:/src/6eWikiScraper, indexed {}".format(W.generated_at),
    "Generated by .claude/skills/bsdata-cat/scripts/wiki_xref.py",
]
for i, t in enumerate(notes, start=2):
    lg.cell(i, 1, t).font = Font(name=ARIAL, size=10,
                                 bold=t.endswith("meanings:") or t.startswith("Fill in"))
lg.column_dimensions["A"].width = 100

base = len(notes) + 3
lg.cell(base, 1, "Counts by category").font = Font(name=ARIAL, bold=True)
for i, cat in enumerate(order, start=base + 1):
    lg.cell(i, 1, cat).font = Font(name=ARIAL, size=10)
    lg.cell(i, 2, '=COUNTIF(Decisions!$A:$A,A{})'.format(i)).font = Font(name=ARIAL, size=10)
undec = base + 5
lg.cell(undec, 1, "Still undecided").font = Font(name=ARIAL, bold=True)
lg.cell(undec, 2, '=COUNTBLANK(Decisions!$I$4:$I${})'.format(len(rows) + 3)).font = Font(
    name=ARIAL, bold=True)
lg.column_dimensions["B"].width = 12

ws = wb.create_sheet("Decisions")
head = ["Category", "Name", "Used by", "File", "Issue",
        "Option A", "Option B", "Wiki page", "Decision", "Notes"]
widths = [18, 30, 40, 22, 40, 30, 44, 44, 22, 30]

ws["A1"] = "Edit the yellow Decision column (and Notes). Row 3 is an example - delete it."
ws["A1"].font = Font(name=ARIAL, size=10, italic=True)
ws.merge_cells("A1:J1")

for c, (h, w) in enumerate(zip(head, widths), start=1):
    cell = ws.cell(2, c, h)
    cell.font = Font(name=ARIAL, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="44546A")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(c)].width = w

example = ["Ambiguous match", "(example) Stream of Corruption",
           "shared; used by Plaguebearers, Stream of Corruption", "Forces of Chaos.cat",
           "wiki has 2 candidate pages", "/magic-item/stream-of-corruption",
           "/special-rules/stream-of-corruption",
           WIKI + "/special-rules/stream-of-corruption",
           "Option B", "unit rule, not the gift"]
for c, v in enumerate(example, start=1):
    ws.cell(3, c, v).font = Font(name=ARIAL, size=10, italic=True, color="808080")

YELLOW = PatternFill("solid", fgColor="FFFF00")
for i, r in enumerate(rows, start=4):
    for c, v in enumerate(r[:8], start=1):
        cell = ws.cell(i, c, v)
        cell.font = Font(name=ARIAL, size=10)
        cell.alignment = Alignment(vertical="top", wrap_text=(c in (3, 5, 6, 7)))
    ws.cell(i, 8).font = Font(name=ARIAL, size=10, color="0563C1", underline="single")
    if r[7]:
        ws.cell(i, 8).hyperlink = r[7]
    ws.cell(i, 9).fill = YELLOW
    ws.cell(i, 10).fill = YELLOW
    carried_d = r[8] if len(r) > 8 else None
    if carried_d:
        ws.cell(i, 9, carried_d["decision"]).font = Font(name=ARIAL, size=10)
        if carried_d.get("note"):
            ws.cell(i, 10, carried_d["note"]).font = Font(name=ARIAL, size=10)

dv = DataValidation(type="list",
                    formula1='"Option A,Option B,Leave as is,Rename catalogue,Needs more thought"',
                    allow_blank=True, showDropDown=False)
ws.add_data_validation(dv)
dv.add("I4:I{}".format(len(rows) + 3))

ws.freeze_panes = "A3"
ws.auto_filter.ref = "A2:J{}".format(len(rows) + 3)

out = REPO / "docs" / "judgement-calls.xlsx"
wb.save(out)
print("wrote {} - {} decision rows".format(out.name, len(rows)))
for k, v in collections.Counter(r[0] for r in rows).most_common():
    print("   {:18} {}".format(k, v))
